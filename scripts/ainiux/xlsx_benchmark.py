#!/usr/bin/env python3
"""Benchmark Ainiux XLSX conversion against PhpSpreadsheet and openpyxl.

Each candidate/input pair runs exactly once in a fresh process with no warmup
or retry. Wall time includes process and library startup, conversion, and
output writes. Neither competitor speaks Markdown natively; this harness owns
a narrow adapter that maps worksheet cell values to GitHub-flavored tables
and Markdown tables back to a new workbook.

Examples:
  python3 scripts/ainiux/xlsx_benchmark.py --self-test
  python3 scripts/ainiux/xlsx_benchmark.py --bootstrap \\
      --ainiux ./ainiux \\
      --phpspreadsheet ~/agent_analysis/PhpSpreadsheet \\
      --output-dir results/xlsx-conversion-YYYY-MM-DD \\
      --report docs/xlsx_conversion_benchmark.md
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import os
import platform
import re
import shlex
import shutil
import subprocess
import sys
import tempfile
import time
import venv
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Mapping, MutableMapping, Optional, Sequence, Tuple
from xml.etree import ElementTree as ET


OPENPYXL_VERSION = "3.1.5"
CANDIDATES = ("ainiux", "openpyxl-adapter", "phpspreadsheet-adapter")
TABLE_RULE_RE = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$")
PHP_WORKER = Path(__file__).resolve().parent / "xlsx_phpspreadsheet_worker.php"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def human_bytes(value: Optional[float]) -> str:
    if value is None:
        return "—"
    units = ("B", "KiB", "MiB", "GiB")
    amount = float(value)
    unit = units[0]
    for unit in units:
        if abs(amount) < 1024.0 or unit == units[-1]:
            break
        amount /= 1024.0
    return f"{amount:.1f} {unit}" if unit != "B" else f"{int(amount)} B"


def candidate_label(name: str) -> str:
    return {
        "ainiux": "Ainiux",
        "openpyxl-adapter": "openpyxl adapter",
        "phpspreadsheet-adapter": "PhpSpreadsheet adapter",
    }.get(name, name)


def stable_environment() -> Dict[str, str]:
    environment = dict(os.environ)
    environment.update({"LC_ALL": "C.UTF-8", "LANG": "C.UTF-8", "PYTHONHASHSEED": "0", "TZ": "UTC"})
    environment.pop("PYTHONPATH", None)
    return environment


def set_affinity(cpu: int) -> None:
    os.sched_setaffinity(0, {cpu})


def run_measured(command: Sequence[str], stdout_path: Path, stderr_path: Path, cpu: int) -> Dict[str, object]:
    import resource

    stdout_path.parent.mkdir(parents=True, exist_ok=True)
    start = time.perf_counter()
    with stdout_path.open("wb") as stdout_stream, stderr_path.open("wb") as stderr_stream:
        process = subprocess.Popen(list(command), stdout=stdout_stream, stderr=stderr_stream,
                                   env=stable_environment(), preexec_fn=lambda: set_affinity(cpu))
        _, status, usage = os.wait4(process.pid, 0)
        process.returncode = os.waitstatus_to_exitcode(status)
    wall = time.perf_counter() - start
    peak_rss_bytes = int(usage.ru_maxrss) * (1 if sys.platform == "darwin" else 1024)
    return {
        "wall_seconds": wall,
        "cpu_seconds": usage.ru_utime + usage.ru_stime,
        "user_cpu_seconds": usage.ru_utime,
        "system_cpu_seconds": usage.ru_stime,
        "peak_rss_bytes": peak_rss_bytes,
        "exit_status": process.returncode,
        "command": list(command),
    }


def git_value(arguments: Sequence[str], cwd: Optional[Path] = None) -> str:
    try:
        completed = subprocess.run(
            ["git"] + list(arguments),
            cwd=str(cwd or Path(__file__).resolve().parents[2]),
            stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, check=True, text=True)
        return completed.stdout.strip()
    except (OSError, subprocess.CalledProcessError):
        return "unknown"


def package_version(distribution: str) -> str:
    try:
        return importlib.metadata.version(distribution)
    except importlib.metadata.PackageNotFoundError:
        return "not installed"


def cpu_description() -> str:
    try:
        text = Path("/proc/cpuinfo").read_text(encoding="utf-8", errors="replace")
    except OSError:
        return platform.processor() or "unknown"
    for line in text.splitlines():
        if line.lower().startswith("model name"):
            return line.split(":", 1)[1].strip()
    return platform.processor() or "unknown"


def escape_gfm_cell(text: str) -> str:
    return text.replace("|", r"\|").replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>")


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        if value.hour or value.minute or value.second:
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return value.strftime("%Y-%m-%d")
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        text = format(value, ".15g")
        return text if text not in ("", "-") else "0"
    return str(value)


def emit_gfm_table(name: str, rows: Sequence[Sequence[str]]) -> str:
    columns = 0
    for row in rows:
        columns = max(columns, len(row))
    if columns == 0:
        return ""
    if columns == 1:
        columns = 2
    lines = ["# " + (name or "Sheet"), ""]

    def emit(row: Sequence[str]) -> str:
        cells = [escape_gfm_cell(row[i]) if i < len(row) else "" for i in range(columns)]
        return "| " + " | ".join(cells) + " |"

    lines.append(emit(rows[0] if rows else []))
    lines.append("|" + " --- |" * columns)
    for row in list(rows)[1:]:
        lines.append(emit(row))
    lines.append("")
    return "\n".join(lines) + "\n"


def split_table_row(line: str) -> List[str]:
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    cells: List[str] = []
    current = []
    escape = False
    for ch in text:
        if escape:
            current.append(ch)
            escape = False
            continue
        if ch == "\\":
            escape = True
            continue
        if ch == "|":
            cells.append("".join(current).strip().replace("<br>", "\n"))
            current = []
            continue
        current.append(ch)
    cells.append("".join(current).strip().replace("<br>", "\n"))
    return cells


def parse_markdown_tables(text: str) -> List[Tuple[str, List[List[str]]]]:
    lines = text.splitlines()
    tables: List[Tuple[str, List[List[str]]]] = []
    pending = "Sheet"
    index = 0
    while index < len(lines):
        stripped = lines[index].strip()
        if stripped.startswith("#"):
            pending = stripped.lstrip("#").strip() or "Sheet"
            index += 1
            continue
        if "|" not in lines[index]:
            index += 1
            continue
        header = split_table_row(lines[index])
        if index + 1 >= len(lines) or not TABLE_RULE_RE.match(lines[index + 1]):
            index += 1
            continue
        rows = [header]
        index += 2
        while index < len(lines) and "|" in lines[index] and lines[index].strip():
            if TABLE_RULE_RE.match(lines[index]):
                index += 1
                continue
            rows.append(split_table_row(lines[index]))
            index += 1
        tables.append((pending, rows))
        pending = "Sheet"
    return tables


def trim_sheet_rows(rows: List[List[str]]) -> List[List[str]]:
    while rows and all(cell == "" for cell in rows[-1]):
        rows.pop()
    width = 0
    for row in rows:
        for index in range(len(row) - 1, -1, -1):
            if row[index] != "":
                width = max(width, index + 1)
                break
    if width == 0:
        return []
    return [row[:width] + [""] * (width - len(row)) if len(row) < width else row[:width] for row in rows]


def openpyxl_to_markdown(input_path: Path, output_path: Path) -> None:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(input_path), data_only=True, read_only=True)
    chunks: List[str] = []
    try:
        for sheet in workbook.worksheets:
            rows: List[List[str]] = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([cell_text(value) for value in row])
            rows = trim_sheet_rows(rows)
            if not rows:
                continue
            chunks.append(emit_gfm_table(sheet.title or "Sheet", rows))
    finally:
        workbook.close()
    if not chunks:
        raise RuntimeError("openpyxl produced empty Markdown")
    output_path.write_text("".join(chunks), encoding="utf-8")


def openpyxl_from_markdown(input_path: Path, output_path: Path) -> None:
    from openpyxl import Workbook

    tables = parse_markdown_tables(input_path.read_text(encoding="utf-8"))
    if not tables:
        raise RuntimeError("Markdown does not contain a valid table for XLSX output")
    workbook = Workbook()
    default = workbook.active
    workbook.remove(default)
    for index, (name, rows) in enumerate(tables):
        title = (name or f"Sheet{index + 1}")[:31]
        sheet = workbook.create_sheet(title=title)
        for row_index, row in enumerate(rows, start=1):
            for column_index, value in enumerate(row, start=1):
                sheet.cell(row_index, column_index, value)
    workbook.save(str(output_path))


def worker_main(candidate: str, direction: str, input_name: str, output_name: str) -> int:
    input_path = Path(input_name)
    output_path = Path(output_name)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if candidate != "openpyxl-adapter":
        raise ValueError("unknown Python worker candidate: " + candidate)
    if direction == "xlsx_to_md":
        openpyxl_to_markdown(input_path, output_path)
    elif direction == "md_to_xlsx":
        openpyxl_from_markdown(input_path, output_path)
    else:
        raise ValueError("unknown direction: " + direction)
    return 0


def read_utf8(path: Path) -> Tuple[Optional[str], Optional[str]]:
    try:
        return path.read_bytes().decode("utf-8"), None
    except (OSError, UnicodeDecodeError) as exc:
        return None, str(exc)


def xlsx_package_validation(path: Path) -> Dict[str, object]:
    try:
        with zipfile.ZipFile(path) as archive:
            bad_member = archive.testzip()
            names = archive.namelist()
            rels = ET.fromstring(archive.read("_rels/.rels"))
            workbook = None
            for child in rels:
                rel_type = child.attrib.get("Type", "")
                if rel_type.endswith("/officeDocument"):
                    workbook = child.attrib.get("Target", "").lstrip("/")
                    break
            if not workbook:
                return {"valid": False, "error": "missing workbook relationship"}
            ET.fromstring(archive.read(workbook))
            return {"valid": bad_member is None, "bad_crc_member": bad_member,
                    "workbook": workbook, "members": len(names)}
    except (OSError, KeyError, ValueError, zipfile.BadZipFile, ET.ParseError) as exc:
        return {"valid": False, "error": str(exc)}


def warnings_from_file(path: Path) -> List[str]:
    text, error = read_utf8(path)
    if error:
        return [error]
    assert text is not None
    return [line for line in text.splitlines() if line.strip()]


def conversion_command(candidate: str, direction: str, input_path: Path, output_path: Path,
                       ainiux: Path, php: Optional[str], autoload: Optional[Path]) -> List[str]:
    if candidate == "ainiux":
        output_format = "md" if direction == "xlsx_to_md" else "xlsx"
        maximum = max(32 * 1024 * 1024, input_path.stat().st_size + 1024 * 1024)
        return [str(ainiux), "--no-config", "--input", str(input_path), "--output-format", output_format,
                "--output", str(output_path), "--max-input-bytes", str(maximum)]
    if candidate == "openpyxl-adapter":
        return [sys.executable, str(Path(__file__).resolve()), "--worker", candidate, direction,
                str(input_path), str(output_path)]
    if candidate == "phpspreadsheet-adapter":
        if not php or autoload is None:
            raise ValueError("PhpSpreadsheet adapter requires php and vendor/autoload.php")
        return [php, str(PHP_WORKER), str(autoload), direction, str(input_path), str(output_path)]
    raise ValueError("unknown candidate: " + candidate)


def default_xlsx_files(repository: Path) -> List[Path]:
    root = repository / "tests/xlsx_files"
    return [
        root / "1mb.xlsx",
        root / "countries-population.xlsx",
        root / "hello-world.xlsx",
        root / "threesheets.xlsx",
        root / "1900_Calendar.xlsx",
        root / "sharedformulae.xlsx",
        root / "utf16be.bom.xlsx",
        root / "without_cell_reference.xlsx",
    ]


def sibling_markdown(xlsx_path: Path) -> Optional[Path]:
    candidate = xlsx_path.with_suffix(".md")
    return candidate if candidate.is_file() else None


def prepare_markdown_sources(xlsx_files: Sequence[Path], ainiux: Path, prepared_dir: Path,
                             cpu: int) -> List[Path]:
    prepared_dir.mkdir(parents=True, exist_ok=True)
    sources: List[Path] = []
    for xlsx_path in xlsx_files:
        sibling = sibling_markdown(xlsx_path)
        if sibling is not None:
            sources.append(sibling)
            continue
        output_path = prepared_dir / (xlsx_path.stem + ".md")
        command = conversion_command("ainiux", "xlsx_to_md", xlsx_path, output_path, ainiux, None, None)
        stdout_path = prepared_dir / (xlsx_path.stem + ".prepare.stdout")
        stderr_path = prepared_dir / (xlsx_path.stem + ".prepare.stderr")
        timing = run_measured(command, stdout_path, stderr_path, cpu)
        if timing["exit_status"] != 0 or not output_path.is_file() or output_path.stat().st_size == 0:
            raise RuntimeError("failed to prepare Markdown from " + str(xlsx_path)
                               + " (exit " + str(timing["exit_status"]) + ")")
        sources.append(output_path)
    return sources


def display_path(path: Path, repository: Path) -> str:
    try:
        return str(path.relative_to(repository))
    except ValueError:
        return str(path).replace(str(Path.home()), "~", 1)


def completeness(timing: Mapping[str, object], output_path: Path, direction: str,
                 parse_error: Optional[str], package: Optional[Mapping[str, object]]) -> List[str]:
    reasons = []
    if timing["exit_status"] != 0:
        reasons.append("exit status " + str(timing["exit_status"]))
    if not output_path.is_file():
        reasons.append("missing output")
    elif output_path.stat().st_size == 0:
        reasons.append("empty output")
    if parse_error:
        reasons.append("format error: " + parse_error)
    if direction == "md_to_xlsx" and package and not package.get("valid"):
        reasons.append("invalid XLSX package")
    return reasons


def aggregate_results(results: Sequence[Mapping[str, object]]) -> Dict[str, object]:
    by_case: Dict[str, set] = {}
    for row in results:
        if row["complete"]:
            by_case.setdefault(str(row["case_id"]), set()).add(str(row["candidate"]))
    all_cases = sorted({str(row["case_id"]) for row in results})
    common = sorted(case for case, names in by_case.items() if names == set(CANDIDATES))

    def totals(case_ids: Sequence[str]) -> List[Dict[str, object]]:
        rows = []
        ainiux_wall = None
        for candidate in CANDIDATES:
            matched = [row for row in results
                       if row["candidate"] == candidate and row["case_id"] in case_ids and row["complete"]]
            wall = sum(float(row["timing"]["wall_seconds"]) for row in matched)  # type: ignore[index]
            cpu = sum(float(row["timing"]["cpu_seconds"]) for row in matched)  # type: ignore[index]
            nbytes = sum(int(row["input_bytes"]) for row in matched)
            if candidate == "ainiux":
                ainiux_wall = wall
            ratio = (wall / ainiux_wall) if ainiux_wall else None
            rows.append({
                "candidate": candidate,
                "cases": len(matched),
                "wall_seconds": wall,
                "cpu_seconds": cpu,
                "input_bytes": nbytes,
                "throughput_bytes_per_second": nbytes / wall if wall else None,
                "wall_ratio_vs_ainiux": ratio,
            })
        return rows

    full_eligible = [candidate for candidate in CANDIDATES
                     if sum(1 for row in results if row["candidate"] == candidate and row["complete"])
                     == len(all_cases)]
    return {
        "common_complete_case_ids": common,
        "common_complete_totals": totals(common),
        "full_corpus_case_ids": all_cases,
        "full_corpus_eligible_candidates": full_eligible,
        "full_corpus_totals": totals(all_cases) if full_eligible else [],
    }


def render_report(payload: Mapping[str, object]) -> str:
    metadata = payload["metadata"]  # type: ignore[assignment]
    results = payload["results"]  # type: ignore[assignment]
    aggregates = payload["aggregates"]  # type: ignore[assignment]
    lines = [
        "# XLSX conversion benchmark (" + str(metadata["date"]) + ")", "",
        "This point-in-time benchmark compares **speed** of XLSX↔Markdown conversion, not lossless "
        "workbook layout. Each of the " + str(len(results)) + " candidate/file combinations ran exactly "
        "once in a fresh process with no warmup or retry. Wall time includes process and library startup, "
        "conversion, and output writes.", "",
        "## Method", "",
        "All conversions were pinned to logical CPU " + str(metadata["cpu_affinity"]) +
        " with `LC_ALL=C.UTF-8`, `PYTHONHASHSEED=0`, and isolated output directories. Candidate order "
        "rotates by input to reduce fixed cache-order bias. Ainiux was invoked by explicit repository "
        "path; its revision and binary digest are recorded below.", "",
        "Neither openpyxl nor PhpSpreadsheet converts Markdown natively. This harness maps worksheet "
        "cell values to GitHub-flavored tables (sheet name as a heading) and Markdown tables back to a "
        "new workbook, matching Ainiux's documented XLSX subset. Formula cells use cached values; the "
        "Excel calculation engine is not run. Drawings, charts, styles, and merged-cell layout are out "
        "of scope.", "",
        "XLSX→Markdown uses the eight in-tree workbooks listed below. Markdown→XLSX uses the adjacent "
        "`countries-population.md` and `hello-world.md` goldens, and Ainiux-extracted Markdown (untimed) "
        "for the other six workbooks so every writer sees the same source.", "",
        "A result is incomplete when conversion fails, the output is missing/empty, Markdown is not "
        "valid UTF-8, or a generated XLSX package has no workbook part. Incomplete results remain "
        "visible but cannot contribute to speed rankings.", "",
        "## Per-file results", "",
        "| Direction | File | Candidate | Input→output | Wall | CPU | Peak RSS | Throughput | Exit | Complete | vs Ainiux |",
        "| --- | --- | --- | ---: | ---: | ---: | ---: | ---: | ---: | --- | ---: |",
    ]
    complete_baselines = {row["case_id"]: row["complete"] for row in results if row["candidate"] == "ainiux"}
    for row in results:
        timing = row["timing"]
        ratio = row.get("wall_ratio_vs_ainiux")
        ratio_eligible = row["complete"] and complete_baselines.get(row["case_id"], False)
        ratio_text = f"{float(ratio):.2f}×" if ratio is not None and ratio_eligible else "—"
        direction = "XLSX→MD" if row["direction"] == "xlsx_to_md" else "MD→XLSX"
        lines.append("| " + " | ".join([
            direction, "`" + str(row["input_display"]).replace("|", r"\|") + "`",
            candidate_label(str(row["candidate"])),
            human_bytes(float(row["input_bytes"])) + "→" + human_bytes(float(row["output_bytes"])),
            f"{float(timing['wall_seconds']):.4f} s",
            f"{float(timing['cpu_seconds']):.4f} s",
            human_bytes(float(timing["peak_rss_bytes"])),
            human_bytes(float(row["throughput_bytes_per_second"])) + "/s",
            str(timing["exit_status"]),
            "yes" if row["complete"] else "**no**",
            ratio_text,
        ]) + " |")

    lines.extend(["", "`vs Ainiux` is candidate wall time divided by Ainiux wall time on the same file; larger is slower.",
                  "", "## Aggregate comparisons", ""])
    common_ids = aggregates["common_complete_case_ids"]
    lines.append("Three-way totals use only the common complete subset: **" + str(len(common_ids)) + " of "
                 + str(len(aggregates["full_corpus_case_ids"])) + " cases**.")
    lines.extend(["", "| Candidate | Eligible cases | Wall total | CPU total | Throughput | vs Ainiux |",
                  "| --- | ---: | ---: | ---: | ---: | ---: |"])
    for row in aggregates["common_complete_totals"]:
        lines.append("| " + " | ".join([
            candidate_label(str(row["candidate"])), str(row["cases"]),
            f"{float(row['wall_seconds']):.4f} s", f"{float(row['cpu_seconds']):.4f} s",
            human_bytes(row["throughput_bytes_per_second"]) + "/s" if row["throughput_bytes_per_second"] else "—",
            f"{float(row['wall_ratio_vs_ainiux']):.2f}×" if row["wall_ratio_vs_ainiux"] is not None else "—",
        ]) + " |")
    eligible = aggregates["full_corpus_eligible_candidates"]
    lines.extend(["", "Full-corpus totals are shown only for candidates completing every case. Eligible: "
                  + (", ".join(candidate_label(item) for item in eligible) if eligible else "none") + ".", ""])
    if aggregates["full_corpus_totals"]:
        lines.extend(["| Candidate | Cases | Wall total | CPU total | Throughput | vs Ainiux |",
                      "| --- | ---: | ---: | ---: | ---: | ---: |"])
        for row in aggregates["full_corpus_totals"]:
            lines.append("| " + " | ".join([
                candidate_label(str(row["candidate"])), str(row["cases"]),
                f"{float(row['wall_seconds']):.4f} s", f"{float(row['cpu_seconds']):.4f} s",
                human_bytes(row["throughput_bytes_per_second"]) + "/s" if row["throughput_bytes_per_second"] else "—",
                f"{float(row['wall_ratio_vs_ainiux']):.2f}×" if row["wall_ratio_vs_ainiux"] is not None else "—",
            ]) + " |")

    complete_counts = {candidate: sum(1 for row in results if row["candidate"] == candidate and row["complete"])
                       for candidate in CANDIDATES}
    common_totals = aggregates["common_complete_totals"]
    fastest = min(common_totals, key=lambda item: item["wall_seconds"]) if common_totals else None
    lines.extend(["", "## Findings", ""])
    if fastest is not None:
        others = "; ".join(
            candidate_label(str(row["candidate"])) + " was " + f"{float(row['wall_ratio_vs_ainiux']):.2f}× Ainiux time"
            for row in common_totals if row["candidate"] != "ainiux")
        lines.append("- On the " + str(len(common_ids)) + "-case common complete subset, **"
                     + candidate_label(str(fastest["candidate"])) + " was fastest** at "
                     + f"{float(fastest['wall_seconds']):.4f} seconds total. " + others + ".")
    lines.append("- Complete conversions: " + ", ".join(
        candidate_label(candidate) + " " + str(complete_counts[candidate]) + "/"
        + str(len(results) // len(CANDIDATES)) for candidate in CANDIDATES) + ".")

    incomplete = [row for row in results if not row["complete"]]
    lines.extend(["", "## Completeness and validation", ""])
    if incomplete:
        lines.extend(["| File | Candidate | Reason |", "| --- | --- | --- |"])
        for row in incomplete:
            reasons = ", ".join(row["incomplete_reasons"])
            lines.append("| `" + str(row["input_display"]).replace("|", r"\|") + "` | "
                         + candidate_label(str(row["candidate"])) + " | " + reasons.replace("|", r"\|") + " |")
    else:
        lines.append("All " + str(len(results)) + " conversions met the completeness rule. Generated XLSX "
                     "packages were checked as ZIP archives with a workbook relationship.")

    warning_rows = [row for row in results if row["warnings"]]
    if warning_rows:
        lines.extend(["", "### Converter diagnostics", "",
                      "These stderr diagnostics are also preserved verbatim beside each output and in `raw.json`.", "",
                      "| File | Candidate | Diagnostic |", "| --- | --- | --- |"])
        for row in warning_rows:
            for warning in row["warnings"]:
                safe_warning = str(warning).replace("|", r"\|").replace("\n", " ")
                lines.append("| `" + str(row["input_display"]).replace("|", r"\|") + "` | "
                             + candidate_label(str(row["candidate"])) + " | " + safe_warning + " |")

    packages = metadata.get("packages", {})
    lines.extend(["", "## Environment", "", "| Item | Value |", "| --- | --- |",
                  "| Timestamp | `" + str(metadata["generated_at_utc"]) + "` |",
                  "| Host | `" + str(metadata["platform"]).replace("|", r"\|") + "` |",
                  "| CPU | `" + str(metadata["cpu"]).replace("|", r"\|") + "` |",
                  "| Python | `" + str(metadata["python"]).replace("|", r"\|") + "` |",
                  "| PHP | `" + str(metadata.get("php", "not recorded")).replace("|", r"\|") + "` |",
                  "| openpyxl | `" + str(packages.get("openpyxl", "not recorded")) + "` |",
                  "| PhpSpreadsheet | `" + str(packages.get("phpspreadsheet", "not recorded")).replace("|", r"\|") + "` |",
                  "| Ainiux revision | `" + str(metadata["ainiux_git_revision"]) + "` |",
                  "| Ainiux SHA-256 | `" + str(metadata["ainiux_sha256"]) + "` |",
                  "| Ainiux binary | `" + str(metadata["ainiux_binary"]) + "` |",
                  "| Harness SHA-256 | `" + str(metadata.get("harness_sha256", "not recorded")) + "` |", "",
                  "Exact harness command:", "", "```sh", str(metadata["exact_command"]), "```", "",
                  "## Limitations and follow-up", "",
                  "These are single observations, not statistically stable latency estimates. OS cache state still "
                  "matters despite fresh processes and rotated order. Competitor numbers belong to this benchmark's "
                  "Markdown adapters because the libraries read and write workbooks but do not provide Markdown "
                  "converters. Adapter mapping of dates, empty cells, and cached formulas can differ from Ainiux "
                  "without failing the speed completeness rule.", "",
                  "This benchmark does not change the production converter. Parser/writer follow-ups should receive "
                  "focused XLSX fixtures and failure-path tests in a separate change.", ""])
    return "\n".join(lines)


def resolve_phpspreadsheet(path: Path) -> Tuple[Path, Path]:
    root = path.expanduser().resolve()
    autoload = root / "vendor" / "autoload.php"
    if not autoload.is_file():
        raise ValueError("PhpSpreadsheet vendor/autoload.php is missing; run --bootstrap or composer install in "
                         + str(root))
    return root, autoload


def php_version(php: str) -> str:
    try:
        completed = subprocess.run([php, "-v"], stdout=subprocess.PIPE, stderr=subprocess.DEVNULL,
                                   check=True, text=True)
        return completed.stdout.splitlines()[0].strip()
    except (OSError, subprocess.CalledProcessError, IndexError):
        return "unknown"


def composer_command(php: str, scratch: Path) -> List[str]:
    composer = shutil.which("composer")
    if composer is not None:
        return [composer]
    phar = scratch / "composer.phar"
    if not phar.is_file():
        installer = scratch / "composer-setup.php"
        completed = subprocess.run(
            [php, "-r",
             "copy('https://getcomposer.org/installer', " + json.dumps(str(installer)) + ");"],
            check=False)
        if completed.returncode != 0:
            raise ValueError("could not download the Composer installer")
        completed = subprocess.run([php, str(installer), "--install-dir", str(scratch),
                                    "--filename", "composer.phar"], check=False)
        if completed.returncode != 0 or not phar.is_file():
            raise ValueError("could not install composer.phar")
    return [php, str(phar)]


def composer_install(phpspreadsheet: Path) -> int:
    if (phpspreadsheet / "vendor" / "autoload.php").is_file():
        return 0
    php = shutil.which("php")
    if php is None:
        raise ValueError("php is required to install PhpSpreadsheet vendor/")
    with tempfile.TemporaryDirectory(prefix="ainiux-xlsx-composer-") as scratch:
        command = composer_command(php, Path(scratch))
        return subprocess.run(command + ["install", "--no-dev", "--no-interaction",
                                         "--ignore-platform-reqs", "--working-dir",
                                         str(phpspreadsheet)], check=False).returncode


def run_one(case_id: str, direction: str, input_path: Path, candidate: str, order_index: int,
            outputs_root: Path, repository: Path, ainiux: Path, php: Optional[str],
            autoload: Optional[Path], cpu: int) -> MutableMapping[str, object]:
    candidate_dir = outputs_root / case_id / candidate
    candidate_dir.mkdir(parents=True)
    output_path = candidate_dir / ("output.md" if direction == "xlsx_to_md" else "output.xlsx")
    stdout_path = candidate_dir / "stdout.txt"
    stderr_path = candidate_dir / "stderr.txt"
    command = conversion_command(candidate, direction, input_path, output_path, ainiux, php, autoload)
    print(f"{input_path.name}: {candidate} ({direction})", file=sys.stderr, flush=True)
    timing = run_measured(command, stdout_path, stderr_path, cpu)
    output_exists = output_path.is_file()
    output_bytes = output_path.stat().st_size if output_exists else 0
    parse_error: Optional[str] = None
    package: Optional[Dict[str, object]] = None
    if output_exists and timing["exit_status"] == 0:
        if direction == "xlsx_to_md":
            _, utf8_error = read_utf8(output_path)
            parse_error = utf8_error
        else:
            package = xlsx_package_validation(output_path)
            if not package.get("valid"):
                parse_error = str(package.get("error", "invalid XLSX package"))
    reasons = completeness(timing, output_path, direction, parse_error, package)
    wall = float(timing["wall_seconds"])
    return {
        "case_id": case_id,
        "input": str(input_path),
        "input_display": display_path(input_path, repository),
        "input_sha256": sha256_file(input_path),
        "input_bytes": input_path.stat().st_size,
        "direction": direction,
        "candidate": candidate,
        "candidate_order": order_index + 1,
        "output": str(output_path),
        "output_bytes": output_bytes,
        "throughput_bytes_per_second": input_path.stat().st_size / wall if wall else 0.0,
        "timing": timing,
        "validation": {"xlsx_package": package} if package is not None else {},
        "warnings": warnings_from_file(stderr_path),
        "complete": not reasons,
        "incomplete_reasons": reasons,
    }


def run_benchmark(args: argparse.Namespace) -> int:
    repository = Path(__file__).resolve().parents[2]
    ainiux = Path(args.ainiux).expanduser().resolve()
    if not ainiux.is_file() or not os.access(ainiux, os.X_OK):
        raise ValueError("Ainiux must be an executable file: " + str(ainiux))
    if not hasattr(os, "sched_setaffinity") or not hasattr(os, "wait4"):
        raise ValueError("fresh-process CPU/RSS measurement requires POSIX sched_setaffinity and wait4")
    allowed = os.sched_getaffinity(0)
    if args.cpu not in allowed:
        raise ValueError(f"logical CPU {args.cpu} is unavailable; allowed CPUs: {sorted(allowed)}")
    if package_version("openpyxl") != OPENPYXL_VERSION:
        raise ValueError("missing pinned openpyxl==" + OPENPYXL_VERSION + " (found "
                         + package_version("openpyxl") + "); use --bootstrap")
    php = str(Path(args.php).expanduser().resolve()) if args.php else shutil.which("php")
    if php is None:
        raise ValueError("php is required for the PhpSpreadsheet adapter")
    phpspreadsheet, autoload = resolve_phpspreadsheet(Path(args.phpspreadsheet))
    if not PHP_WORKER.is_file():
        raise ValueError("missing PHP worker: " + str(PHP_WORKER))

    xlsx_files = [Path(item).expanduser().resolve() for item in args.inputs] if args.inputs else default_xlsx_files(repository)
    if not xlsx_files or len(xlsx_files) > 8:
        raise ValueError("benchmark accepts 1-8 XLSX sources, got " + str(len(xlsx_files)))
    missing = [str(path) for path in xlsx_files if not path.is_file()]
    if missing:
        raise ValueError("missing benchmark input(s): " + ", ".join(missing))
    if len(set(xlsx_files)) != len(xlsx_files):
        raise ValueError("benchmark inputs must be unique")
    for path in xlsx_files:
        if path.suffix.casefold() != ".xlsx":
            raise ValueError("XLSX corpus entries must use a .xlsx ending: " + str(path))

    output_root = Path(args.output_dir).expanduser().resolve()
    if output_root.exists() and any(output_root.iterdir()):
        raise ValueError("output directory must be absent or empty (no overwrite): " + str(output_root))
    output_root.mkdir(parents=True, exist_ok=True)
    outputs_root = output_root / "outputs"
    outputs_root.mkdir()

    print("preparing Markdown writer sources", file=sys.stderr, flush=True)
    md_files = prepare_markdown_sources(xlsx_files, ainiux, output_root / "prepared-md", args.cpu)

    jobs: List[Tuple[str, str, Path]] = []
    for index, path in enumerate(xlsx_files):
        jobs.append((f"{index + 1:02d}-xlsx_to_md-{path.stem}", "xlsx_to_md", path))
    for index, path in enumerate(md_files):
        jobs.append((f"{index + 1:02d}-md_to_xlsx-{path.stem}", "md_to_xlsx", path))

    results: List[MutableMapping[str, object]] = []
    total = len(jobs) * len(CANDIDATES)
    for job_index, (case_id, direction, input_path) in enumerate(jobs):
        order = list(CANDIDATES[job_index % len(CANDIDATES):] + CANDIDATES[:job_index % len(CANDIDATES)])
        for order_index, candidate in enumerate(order):
            print(f"[{len(results) + 1:02d}/{total:02d}] ", end="", file=sys.stderr, flush=True)
            results.append(run_one(case_id, direction, input_path, candidate, order_index,
                                   outputs_root, repository, ainiux, php, autoload, args.cpu))

    by_case: Dict[str, Dict[str, MutableMapping[str, object]]] = {}
    for row in results:
        by_case.setdefault(str(row["case_id"]), {})[str(row["candidate"])] = row
    for candidate_rows in by_case.values():
        baseline = float(candidate_rows["ainiux"]["timing"]["wall_seconds"])  # type: ignore[index]
        for row in candidate_rows.values():
            row["wall_ratio_vs_ainiux"] = float(row["timing"]["wall_seconds"]) / baseline  # type: ignore[index]

    generated = datetime.now(timezone.utc)
    php_identity = git_value(["rev-parse", "--short", "HEAD"], cwd=phpspreadsheet)
    metadata = {
        "schema_version": 1,
        "date": generated.date().isoformat(),
        "generated_at_utc": generated.isoformat(),
        "platform": platform.platform(),
        "cpu": cpu_description(),
        "cpu_affinity": args.cpu,
        "python": sys.version.replace("\n", " "),
        "php": php_version(php),
        "packages": {
            "openpyxl": package_version("openpyxl"),
            "phpspreadsheet": str(phpspreadsheet) + " @" + php_identity,
        },
        "ainiux_binary": str(ainiux),
        "ainiux_sha256": sha256_file(ainiux),
        "ainiux_git_revision": git_value(["rev-parse", "HEAD"]),
        "harness_sha256": sha256_file(Path(__file__).resolve()),
        "ainiux_git_dirty": bool(git_value(["status", "--porcelain"])),
        "candidate_order_policy": "cyclic rotation by case index",
        "runs_per_pair": 1,
        "warmups": 0,
        "retries": 0,
        "exact_command": os.environ.get("AINIUX_XLSX_BENCH_COMMAND",
                                        shlex.join([sys.executable, str(Path(__file__).resolve())] + sys.argv[1:])),
    }
    payload: Dict[str, object] = {
        "metadata": metadata,
        "inputs": [str(path) for path in xlsx_files],
        "markdown_inputs": [str(path) for path in md_files],
        "results": results,
        "aggregates": aggregate_results(results),
    }
    raw_path = output_root / "raw.json"
    raw_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    report_text = render_report(payload)
    generated_report = output_root / "report.md"
    generated_report.write_text(report_text, encoding="utf-8")
    if args.report:
        report_path = Path(args.report).expanduser().resolve()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(report_text, encoding="utf-8")
    print(str(raw_path))
    print(str(generated_report))
    return 0


def self_test() -> int:
    markdown = "# People\n\n| Name | Value |\n| --- | --- |\n| Hello | 7 |\n"
    tables = parse_markdown_tables(markdown)
    assert tables == [("People", [["Name", "Value"], ["Hello", "7"]])], tables
    rendered = emit_gfm_table("People", [["Name", "Value"], ["Hello|x", "7"]])
    assert rendered.startswith("# People\n")
    assert "Hello\\|x" in rendered
    round_trip = parse_markdown_tables(rendered)
    assert round_trip[0][1][1][0] == "Hello|x"
    assert trim_sheet_rows([["a", ""], ["", ""], ["", ""]]) == [["a"]]
    with tempfile.TemporaryDirectory(prefix="ainiux-xlsx-self-test-") as temporary:
        root = Path(temporary)
        md_path = root / "in.md"
        xlsx_path = root / "out.xlsx"
        md_path.write_text(markdown, encoding="utf-8")
        try:
            import openpyxl  # noqa: F401
        except ImportError:
            package = xlsx_package_validation(Path(__file__))
            assert package["valid"] is False
        else:
            openpyxl_from_markdown(md_path, xlsx_path)
            package = xlsx_package_validation(xlsx_path)
            assert package["valid"], package
            openpyxl_to_markdown(xlsx_path, root / "back.md")
            back = (root / "back.md").read_text(encoding="utf-8")
            assert "Hello" in back and "People" in back
    sample_payload = {
        "metadata": {"date": "2000-01-01", "generated_at_utc": "2000-01-01T00:00:00+00:00",
                     "cpu_affinity": 5, "platform": "test", "cpu": "test", "python": "test",
                     "php": "test", "packages": {"openpyxl": OPENPYXL_VERSION, "phpspreadsheet": "test"},
                     "ainiux_git_revision": "deadbeef", "ainiux_sha256": "0" * 64,
                     "ainiux_binary": "./ainiux", "exact_command": "python3 xlsx_benchmark.py"},
        "results": [],
        "aggregates": {"common_complete_case_ids": [], "common_complete_totals": [],
                       "full_corpus_case_ids": [], "full_corpus_eligible_candidates": [],
                       "full_corpus_totals": []},
    }
    report = render_report(sample_payload)
    assert "XLSX conversion benchmark" in report and "Aggregate comparisons" in report
    print("XLSX benchmark self-test passed")
    return 0


def bootstrap(argv: Sequence[str]) -> int:
    forwarded = [item for item in argv if item != "--bootstrap"]
    exact = shlex.join([sys.executable, str(Path(__file__).resolve())] + list(argv))
    args = parse_args(forwarded)
    phpspreadsheet = Path(args.phpspreadsheet).expanduser().resolve()
    if not phpspreadsheet.is_dir():
        raise ValueError("PhpSpreadsheet tree is missing: " + str(phpspreadsheet))
    status = composer_install(phpspreadsheet)
    if status != 0:
        return status
    with tempfile.TemporaryDirectory(prefix="ainiux-xlsx-benchmark-venv-") as temporary:
        environment_path = Path(temporary)
        venv.EnvBuilder(with_pip=True, clear=True).create(environment_path)
        python = environment_path / ("Scripts/python.exe" if os.name == "nt" else "bin/python")
        install = [str(python), "-m", "pip", "install", "--disable-pip-version-check",
                   "openpyxl==" + OPENPYXL_VERSION]
        completed = subprocess.run(install, check=False)
        if completed.returncode != 0:
            return completed.returncode
        environment = dict(os.environ)
        environment["AINIUX_XLSX_BENCH_COMMAND"] = exact
        return subprocess.run([str(python), str(Path(__file__).resolve())] + forwarded,
                              check=False, env=environment).returncode


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("inputs", nargs="*", help="up to eight .xlsx sources (default: documented eight-file matrix)")
    parser.add_argument("--ainiux", default="./ainiux", help="explicit repository binary (default: ./ainiux)")
    parser.add_argument("--phpspreadsheet", default=str(Path.home() / "agent_analysis/PhpSpreadsheet"),
                        help="PhpSpreadsheet source tree (needs vendor/autoload.php)")
    parser.add_argument("--php", help="php CLI used for the PhpSpreadsheet adapter (default: php on PATH)")
    parser.add_argument("--output-dir", default="results/xlsx-conversion-" + datetime.now().date().isoformat())
    parser.add_argument("--report", help="also write the generated Markdown report here")
    parser.add_argument("--cpu", type=int, default=5, help="logical CPU affinity for every candidate process")
    parser.add_argument("--bootstrap", action="store_true",
                        help="composer-install PhpSpreadsheet and run inside a temporary venv with pinned openpyxl")
    parser.add_argument("--self-test", action="store_true", help="run synthetic table parse/emit/report tests")
    parser.add_argument("--worker", nargs=4, metavar=("CANDIDATE", "DIRECTION", "INPUT", "OUTPUT"),
                        help=argparse.SUPPRESS)
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    actual = list(argv) if argv is not None else sys.argv[1:]
    args = parse_args(actual)
    try:
        if args.worker:
            return worker_main(*args.worker)
        if args.self_test:
            return self_test()
        if args.bootstrap:
            return bootstrap(actual)
        return run_benchmark(args)
    except (OSError, ValueError, RuntimeError, zipfile.BadZipFile, ET.ParseError) as exc:
        print("xlsx benchmark: " + str(exc), file=sys.stderr)
        return 2


if __name__ == "__main__":
    sys.exit(main())
